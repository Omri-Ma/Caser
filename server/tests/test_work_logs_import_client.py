from io import BytesIO

from openpyxl import Workbook

from conftest import auth_for, make_assignment, make_case, make_identity, make_membership, make_tenant
from shared.models import WorkLog
from shared.models.enums import CaseStatus, UserRole, WorkLogSource

HEADER = ["תיק", "תאריך (YYYY-MM-DD)", "שעות", "תיאור"]


def _lawyer_on_case(db, tenant, case, email="lawyer@acme.com", name="Lior Lawyer"):
    identity = make_identity(db, email, name)
    membership = make_membership(db, identity.id, tenant.id, UserRole.LAWYER)
    make_assignment(db, tenant.id, case.id, membership.id)
    return identity, membership


def _build_xlsx(rows: list[list]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Import"
    sheet.append(HEADER)
    for row in rows:
        sheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _upload(client, url, content, headers, cookies, filename="import.xlsx"):
    return client.post(
        url,
        headers=headers,
        cookies=cookies,
        files={"file": (filename, content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )


def test_download_template_lists_only_assigned_open_cases(client_client, db):
    tenant = make_tenant(db, "acme")
    open_case = make_case(db, tenant.id, title="Open Case")
    closed_case = make_case(db, tenant.id, title="Closed Case", status=CaseStatus.CLOSED)
    lawyer_identity, lawyer_membership = _lawyer_on_case(db, tenant, open_case)
    make_assignment(db, tenant.id, closed_case.id, lawyer_membership.id)
    headers, cookies = auth_for(lawyer_identity, "acme")

    resp = client_client.get("/work-logs/import/template", headers=headers, cookies=cookies)

    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith("application/vnd.openxmlformats")

    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(resp.content))
    options = [cell[0] for cell in workbook["Cases"].iter_rows(values_only=True)]
    assert any("Open Case" in opt for opt in options)
    assert not any("Closed Case" in opt for opt in options)


def test_self_import_creates_work_logs_with_excel_source(client_client, db):
    tenant = make_tenant(db, "acme")
    case = make_case(db, tenant.id, title="Case A")
    lawyer_identity, _ = _lawyer_on_case(db, tenant, case)
    headers, cookies = auth_for(lawyer_identity, "acme")

    content = _build_xlsx(
        [
            [f"{case.id} - Case A", "2026-09-01", "3.5", "Drafted motion"],
            [f"{case.id} - Case A", "2026-09-02", "2", None],
        ]
    )

    resp = _upload(client_client, "/work-logs/import", content, headers, cookies)

    assert resp.status_code == 201
    assert resp.json()["imported_count"] == 2

    logs = db.query(WorkLog).filter(WorkLog.tenant_id == tenant.id).all()
    assert len(logs) == 2
    assert all(log.source == WorkLogSource.EXCEL_IMPORT for log in logs)


def test_import_rejects_whole_file_on_one_bad_row(client_client, db):
    tenant = make_tenant(db, "acme")
    case = make_case(db, tenant.id, title="Case A")
    lawyer_identity, _ = _lawyer_on_case(db, tenant, case)
    headers, cookies = auth_for(lawyer_identity, "acme")

    content = _build_xlsx(
        [
            [f"{case.id} - Case A", "2026-09-01", "3.5", "Good row"],
            [f"{case.id} - Case A", "not-a-date", "2", "Bad row"],
        ]
    )

    resp = _upload(client_client, "/work-logs/import", content, headers, cookies)

    assert resp.status_code == 422
    body = resp.json()
    assert len(body["row_errors"]) == 1
    assert body["row_errors"][0]["row"] == 3  # header=1, good row=2, bad row=3

    # No partial import — the good row was not persisted either.
    assert db.query(WorkLog).filter(WorkLog.tenant_id == tenant.id).count() == 0


def test_import_rejects_unassigned_case(client_client, db):
    tenant = make_tenant(db, "acme")
    case = make_case(db, tenant.id, title="Case A")
    other_case = make_case(db, tenant.id, title="Not Mine")
    lawyer_identity, _ = _lawyer_on_case(db, tenant, case)
    headers, cookies = auth_for(lawyer_identity, "acme")

    content = _build_xlsx([[f"{other_case.id} - Not Mine", "2026-09-01", "2", None]])

    resp = _upload(client_client, "/work-logs/import", content, headers, cookies)

    assert resp.status_code == 422
    assert "not assigned" in resp.json()["row_errors"][0]["message"]


def test_import_rejects_closed_case(client_client, db):
    tenant = make_tenant(db, "acme")
    case = make_case(db, tenant.id, title="Case A", status=CaseStatus.CLOSED)
    lawyer_identity, _ = _lawyer_on_case(db, tenant, case)
    headers, cookies = auth_for(lawyer_identity, "acme")

    content = _build_xlsx([[f"{case.id} - Case A", "2026-09-01", "2", None]])

    resp = _upload(client_client, "/work-logs/import", content, headers, cookies)

    assert resp.status_code == 422
    assert "closed" in resp.json()["row_errors"][0]["message"]


def test_import_rejects_non_positive_hours(client_client, db):
    tenant = make_tenant(db, "acme")
    case = make_case(db, tenant.id, title="Case A")
    lawyer_identity, _ = _lawyer_on_case(db, tenant, case)
    headers, cookies = auth_for(lawyer_identity, "acme")

    content = _build_xlsx([[f"{case.id} - Case A", "2026-09-01", "0", None]])

    resp = _upload(client_client, "/work-logs/import", content, headers, cookies)

    assert resp.status_code == 422
    assert "positive" in resp.json()["row_errors"][0]["message"]


def test_import_rejects_non_xlsx_file(client_client, db):
    tenant = make_tenant(db, "acme")
    case = make_case(db, tenant.id)
    lawyer_identity, _ = _lawyer_on_case(db, tenant, case)
    headers, cookies = auth_for(lawyer_identity, "acme")

    resp = client_client.post(
        "/work-logs/import",
        headers=headers,
        cookies=cookies,
        files={"file": ("import.txt", b"not really an excel file", "text/plain")},
    )

    assert resp.status_code == 400


def test_client_cannot_import_work_logs(client_client, db):
    tenant = make_tenant(db, "acme")
    case = make_case(db, tenant.id)
    identity = make_identity(db, "client@acme.com", "Dana Client")
    membership = make_membership(db, identity.id, tenant.id, UserRole.CLIENT)
    make_assignment(db, tenant.id, case.id, membership.id)
    headers, cookies = auth_for(identity, "acme")

    content = _build_xlsx([[f"{case.id} - Case", "2026-09-01", "2", None]])
    resp = _upload(client_client, "/work-logs/import", content, headers, cookies)

    assert resp.status_code == 403
