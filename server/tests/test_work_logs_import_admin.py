from io import BytesIO

from openpyxl import Workbook

from conftest import auth_for, make_assignment, make_case, make_identity, make_membership, make_tenant
from shared.models import AuditLog, WorkLog
from shared.models.enums import CaseStatus, UserRole, WorkLogSource

HEADER = ["Lawyer Email", "Case", "Date (YYYY-MM-DD)", "Hours", "Description"]


def _office_manager(db, tenant):
    identity = make_identity(db, "manager@acme.com", "Ofra Manager")
    membership = make_membership(db, identity.id, tenant.id, UserRole.OFFICE_MANAGER)
    return identity, membership


def _lawyer_on_case(db, tenant, case, email="lawyer@acme.com", name="Lior Lawyer"):
    identity = make_identity(db, email, name)
    membership = make_membership(db, identity.id, tenant.id, UserRole.LAWYER)
    make_assignment(db, tenant.id, case.id, membership.id)
    return identity, membership


def _build_xlsx(rows: list[list]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Import"
    sheet.append(HEADER)
    for row in rows:
        sheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _upload(client, content, headers, cookies, filename="import.xlsx"):
    return client.post(
        "/work-logs/import",
        headers=headers,
        cookies=cookies,
        files={"file": (filename, content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )


def test_download_template_lists_every_open_case_and_has_email_column(admin_client, db):
    tenant = make_tenant(db, "acme")
    make_case(db, tenant.id, title="Case A")
    make_case(db, tenant.id, title="Case B", status=CaseStatus.CLOSED)
    manager_identity, _ = _office_manager(db, tenant)
    headers, cookies = auth_for(manager_identity, "acme")

    resp = admin_client.get("/work-logs/import/template", headers=headers, cookies=cookies)

    assert resp.status_code == 200
    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(resp.content))
    header_row = next(workbook["Import"].iter_rows(values_only=True))
    assert header_row[0] == "Lawyer Email"

    options = [cell[0] for cell in workbook["Cases"].iter_rows(values_only=True)]
    assert any("Case A" in opt for opt in options)
    assert not any("Case B" in opt for opt in options)


def test_bulk_import_creates_work_logs_and_audits_the_import(admin_client, db):
    tenant = make_tenant(db, "acme")
    case = make_case(db, tenant.id, title="Case A")
    manager_identity, manager_membership = _office_manager(db, tenant)
    lawyer_identity, _ = _lawyer_on_case(db, tenant, case)
    headers, cookies = auth_for(manager_identity, "acme")

    content = _build_xlsx([["lawyer@acme.com", f"{case.id} - Case A", "2026-09-01", "3", "Historical entry"]])

    resp = _upload(admin_client, content, headers, cookies)

    assert resp.status_code == 201
    assert resp.json()["imported_count"] == 1

    log = db.query(WorkLog).filter(WorkLog.tenant_id == tenant.id).first()
    assert log.source == WorkLogSource.EXCEL_IMPORT
    assert log.lawyer_id != manager_membership.id  # hours belong to the lawyer, not the importing manager

    audit = db.query(AuditLog).filter(AuditLog.action == "work_log_excel_imported").first()
    assert audit is not None
    assert audit.user_id == manager_membership.id  # who triggered it, kept separate from whose hours they are


def test_bulk_import_rejects_unknown_lawyer_email(admin_client, db):
    tenant = make_tenant(db, "acme")
    case = make_case(db, tenant.id, title="Case A")
    manager_identity, _ = _office_manager(db, tenant)
    headers, cookies = auth_for(manager_identity, "acme")

    content = _build_xlsx([["nobody@acme.com", f"{case.id} - Case A", "2026-09-01", "2", None]])

    resp = _upload(admin_client, content, headers, cookies)

    assert resp.status_code == 422
    assert "No active lawyer" in resp.json()["row_errors"][0]["message"]


def test_bulk_import_rejects_lawyer_from_another_tenant(admin_client, db):
    tenant_a = make_tenant(db, "acme")
    tenant_b = make_tenant(db, "beta")
    case_a = make_case(db, tenant_a.id, title="Case A")
    manager_identity, _ = _office_manager(db, tenant_a)
    headers, cookies = auth_for(manager_identity, "acme")

    # A real lawyer, but a member of a different tenant entirely.
    other_case = make_case(db, tenant_b.id, title="Other Firm Case")
    _lawyer_on_case(db, tenant_b, other_case, email="cross-tenant@beta.com")

    content = _build_xlsx([["cross-tenant@beta.com", f"{case_a.id} - Case A", "2026-09-01", "2", None]])

    resp = _upload(admin_client, content, headers, cookies)

    assert resp.status_code == 422
    assert "No active lawyer" in resp.json()["row_errors"][0]["message"]


def test_bulk_import_rejects_inactive_lawyer(admin_client, db):
    tenant = make_tenant(db, "acme")
    case = make_case(db, tenant.id, title="Case A")
    manager_identity, _ = _office_manager(db, tenant)
    lawyer_identity, lawyer_membership = _lawyer_on_case(db, tenant, case)
    lawyer_membership.active = False
    db.commit()
    headers, cookies = auth_for(manager_identity, "acme")

    content = _build_xlsx([["lawyer@acme.com", f"{case.id} - Case A", "2026-09-01", "2", None]])

    resp = _upload(admin_client, content, headers, cookies)

    assert resp.status_code == 422
    assert "No active lawyer" in resp.json()["row_errors"][0]["message"]


def test_lawyer_cannot_use_admin_bulk_import(admin_client, db):
    tenant = make_tenant(db, "acme")
    case = make_case(db, tenant.id, title="Case A")
    lawyer_identity, _ = _lawyer_on_case(db, tenant, case)
    headers, cookies = auth_for(lawyer_identity, "acme")

    content = _build_xlsx([["lawyer@acme.com", f"{case.id} - Case A", "2026-09-01", "2", None]])

    resp = _upload(admin_client, content, headers, cookies)

    assert resp.status_code == 403
